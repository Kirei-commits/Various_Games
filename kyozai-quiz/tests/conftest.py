from __future__ import annotations

import io
import shutil
from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches

from kyozai_quiz import db
from kyozai_quiz.config import Settings
from kyozai_quiz.schemas import FigureAnalysis, GeneratedQuestion, QuestionSet

HAS_SOFFICE = shutil.which("soffice") is not None
needs_soffice = pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice がない")


def _png(width=400, height=300) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), 0)
    pix.set_rect(pix.irect, (40, 120, 200))
    return pix.tobytes("png")


@pytest.fixture
def sample_pptx(tmp_path) -> Path:
    """1: テキストのみ / 2: 図形と矢印のフロー / 3: 画像 / 4: グラフ / 5: 表（図なし扱い）/ 6: 非表示・小さいロゴのみ"""
    prs = Presentation()
    blank, title_body = prs.slide_layouts[6], prs.slide_layouts[1]

    s = prs.slides.add_slide(title_body)
    s.shapes.title.text = "情報セキュリティの基本"
    s.placeholders[1].text = "機密性・完全性・可用性の3要素"
    s.notes_slide.notes_text_frame.text = "CIA と呼ばれることを補足する"

    s = prs.slides.add_slide(blank)
    box = s.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(8), Inches(0.8))
    box.text_frame.text = "インシデント対応の流れ"
    a = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(2), Inches(2), Inches(1))
    a.text_frame.text = "検知"
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4), Inches(2), Inches(2), Inches(1))
    b.text_frame.text = "封じ込め"
    s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2.5), Inches(2.5), Inches(4), Inches(2.5))

    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(io.BytesIO(_png()), Inches(1), Inches(1), Inches(6), Inches(4))

    s = prs.slides.add_slide(blank)
    data = CategoryChartData()
    data.categories = ["2023", "2024"]
    data.add_series("件数", (12, 30))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data)

    s = prs.slides.add_slide(blank)
    table = s.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "要素", "意味"
    table.cell(1, 0).text, table.cell(1, 1).text = "機密性", "許可された人だけ"

    s = prs.slides.add_slide(blank)
    s._element.set("show", "0")
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = "非表示スライド"
    s.shapes.add_picture(io.BytesIO(_png(20, 20)), Inches(9), Inches(7), Inches(0.3), Inches(0.3))

    path = tmp_path / "セキュリティ入門.pptx"
    prs.save(path)
    return path


@pytest.fixture
def settings(tmp_path) -> Settings:
    s = Settings.from_env()
    s.data_dir = tmp_path / "data"
    s.gemini_api_key = "test-key"
    s.workers = 1
    s.mermaid_validator = "basic"
    return s


@pytest.fixture
def conn(settings):
    c = db.connect(settings.db_path)
    yield c
    c.close()


class FakeClient:
    """Gemini の代わり。呼ばれた内容を記録し、決められた応答を返す。"""

    def __init__(self, model_name="fake-vlm-1", fail_if=None, mermaid=None, questions=None):
        self.model_name = model_name
        self.fail_if = fail_if or (lambda prompt: False)  # プロンプトを見て失敗させるか決める
        self.mermaid = mermaid if mermaid is not None else ['flowchart LR\n  A["検知"] --> B["封じ込め"]']
        self.questions = questions
        self.calls: list[dict] = []

    def generate_json(self, prompt, schema, images=None):
        self.calls.append({"prompt": prompt, "schema": schema, "images": images or []})
        if self.fail_if(prompt):
            raise RuntimeError("503 overloaded")
        if schema is FigureAnalysis:
            return FigureAnalysis(
                description="「検知」から「封じ込め」へ矢印が伸びている。", mermaid=list(self.mermaid), tables=["| a | b |\n|---|---|\n| 1 | 2 |"]
            )
        if schema is QuestionSet:
            return QuestionSet(questions=self.questions or default_questions())
        raise AssertionError(schema)


def default_questions():
    return [
        GeneratedQuestion(
            question_type="multiple_choice", question="最初に行うのは？", choices=["検知", "封じ込め", "復旧", "報告"],
            answer="A", explanation="図の矢印の始点", figure_based=True,
        ),
        GeneratedQuestion(
            question_type="true_false", question="封じ込めの後に検知を行う。", choices=[], answer="誤",
            explanation="矢印は検知→封じ込め", figure_based=True,
        ),
        GeneratedQuestion(
            question_type="short_answer", question="検知の次の段階は？", choices=[], answer="封じ込め",
            explanation="図より", figure_based=True,
        ),
        GeneratedQuestion(
            question_type="multiple_choice", question="壊れた問題", choices=["a", "b"], answer="a",
            explanation="", figure_based=False,
        ),
    ]
