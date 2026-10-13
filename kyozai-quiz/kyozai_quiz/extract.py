"""python-pptx でスライドのテキスト・発表者ノートを取り出し、図を含むかを判定する（OCR は使わない）。"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture

SMARTART_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "dgm": SMARTART_URI,
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

# 図とみなす種類
KIND_PICTURE = "picture"  # 画像
KIND_CHART = "chart"  # グラフ
KIND_SMARTART = "smartart"
KIND_SHAPE = "shape"  # 図形（オートシェイプ・フリーフォーム・線・コネクタ）
KIND_GROUP = "group"  # グループ化された図形


@dataclass
class SlideContent:
    page_number: int  # 1始まり。PDF のページ番号と一致させる
    text: str
    notes: str
    figure_kinds: list[str] = field(default_factory=list)
    hidden: bool = False

    @property
    def has_figure(self) -> bool:
        return bool(self.figure_kinds)


def extract_slides(pptx_path: Path, min_picture_area_ratio: float = 0.01) -> list[SlideContent]:
    prs = Presentation(str(pptx_path))
    slide_area = (prs.slide_width or 1) * (prs.slide_height or 1)
    result = []
    for index, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        kinds: set[str] = set()
        for shape in _reading_order(slide.shapes):
            _walk(shape, slide, texts, kinds, slide_area, min_picture_area_ratio)
        result.append(
            SlideContent(
                page_number=index,
                text="\n".join(t for t in texts if t.strip()),
                notes=_notes(slide),
                figure_kinds=sorted(kinds),
                hidden=slide._element.get("show") == "0",
            )
        )
    return result


def _reading_order(shapes):
    return sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0)))


def _walk(shape, slide, texts, kinds, slide_area, min_ratio, in_group=False):
    if isinstance(shape, GroupShape):
        kinds.add(KIND_GROUP)
        for child in _reading_order(shape.shapes):
            _walk(child, slide, texts, kinds, slide_area, min_ratio, in_group=True)
        return

    if isinstance(shape, Picture):
        area = (shape.width or 0) * (shape.height or 0)
        if area / slide_area >= min_ratio:
            kinds.add(KIND_PICTURE)
        return

    if getattr(shape, "has_chart", False) and shape.has_chart:
        kinds.add(KIND_CHART)
        texts.append(_chart_text(shape.chart))
        return

    if getattr(shape, "has_table", False) and shape.has_table:
        # 表は python-pptx から構造ごと取れるので、図ではなくテキスト（マークダウン表）として扱う
        texts.append(_table_markdown(shape.table))
        return

    if _is_smartart(shape):
        kinds.add(KIND_SMARTART)
        smartart_text = _smartart_text(shape, slide)
        if smartart_text:
            texts.append(f"[SmartArt] {smartart_text}")
        return

    shape_type = _safe_shape_type(shape)
    is_placeholder = getattr(shape, "is_placeholder", False)
    if shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.LINE) and not is_placeholder:
        kinds.add(KIND_SHAPE)
    elif shape._element.tag.endswith("}cxnSp"):  # コネクタ（矢印線）
        kinds.add(KIND_SHAPE)

    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
        if text:
            texts.append(text)


def _safe_shape_type(shape):
    try:
        return shape.shape_type
    except NotImplementedError:
        return None


def _is_smartart(shape) -> bool:
    uris = shape._element.xpath("./a:graphic/a:graphicData/@uri")
    return SMARTART_URI in uris


def _smartart_text(shape, slide) -> str:
    """SmartArt の文字はスライド本体ではなく diagram data パートにあるので、そこから読む。"""
    rel_ids = [el.get(f"{{{NS['r']}}}dm") for el in shape._element.iter(f"{{{SMARTART_URI}}}relIds")]
    parts = []
    for rel_id in rel_ids:
        if not rel_id:
            continue
        try:
            data = etree.fromstring(slide.part.related_part(rel_id).blob)
        except Exception:  # 壊れた SmartArt はテキストなしとして扱う
            continue
        parts.extend(t.text for t in data.iter(f"{{{NS['a']}}}t") if t.text and t.text.strip())
    return " / ".join(parts)


def _chart_text(chart) -> str:
    lines = []
    try:
        title = chart.chart_title.text_frame.text if chart.has_title else ""
    except Exception:
        title = ""
    lines.append(f"[グラフ] {title}".rstrip())
    try:
        for plot in chart.plots:
            categories = [str(c) for c in plot.categories]
            if categories:
                lines.append("項目: " + ", ".join(categories))
            for series in plot.series:
                values = ", ".join("" if v is None else f"{v:g}" for v in series.values)
                lines.append(f"系列「{series.name}」: {values}")
    except Exception:
        pass
    return "\n".join(lines)


def _table_markdown(table) -> str:
    rows = [[_cell(c.text) for c in row.cells] for row in table.rows]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def _cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", "<br>").strip()


def _notes(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return frame.text.strip() if frame is not None else ""
